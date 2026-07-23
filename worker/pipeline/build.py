"""Build Word letterhead packages from a full-page artwork PNG."""

from __future__ import annotations

import zipfile
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .detect import Detection, MarginsMm


@dataclass
class BuildOptions:
    margins: MarginsMm
    page_width_mm: float
    page_height_mm: float
    font_name: str = "Calibri"
    font_size_pt: float = 11.0
    font_color: Tuple[int, int, int] = (0x1A, 0x1A, 0x1A)


def set_run_font(run, name: str, size_pt: float, color: Optional[RGBColor] = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color


def add_page_background_picture(header_paragraph, image_path: Path, width_mm: float, height_mm: float) -> None:
    """Place a full-page image in the header, page-anchored, behind text."""
    run = header_paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Mm(width_mm), height=Mm(height_mm))

    inline = inline_shape._inline
    drawing = inline.getparent()
    extent = inline.extent
    cx = extent.cx
    cy = extent.cy
    doc_pr = inline.docPr
    graphic = inline.graphic

    anchor = OxmlElement("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", "0")
    anchor.set("behindDoc", "1")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "page")
    pos_h_offset = OxmlElement("wp:posOffset")
    pos_h_offset.text = "0"
    pos_h.append(pos_h_offset)
    anchor.append(pos_h)

    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    pos_v_offset = OxmlElement("wp:posOffset")
    pos_v_offset.text = "0"
    pos_v.append(pos_v_offset)
    anchor.append(pos_v)

    extent_el = OxmlElement("wp:extent")
    extent_el.set("cx", str(cx))
    extent_el.set("cy", str(cy))
    anchor.append(extent_el)

    effect = OxmlElement("wp:effectExtent")
    effect.set("l", "0")
    effect.set("t", "0")
    effect.set("r", "0")
    effect.set("b", "0")
    anchor.append(effect)

    wrap = OxmlElement("wp:wrapNone")
    anchor.append(wrap)
    anchor.append(deepcopy(doc_pr))

    c_nv = inline.find(qn("wp:cNvGraphicFramePr"))
    if c_nv is not None:
        anchor.append(deepcopy(c_nv))
    else:
        c_nv_new = OxmlElement("wp:cNvGraphicFramePr")
        graphic_locks = OxmlElement("a:graphicFrameLocks")
        graphic_locks.set(qn("xmlns:a"), "http://schemas.openxmlformats.org/drawingml/2006/main")
        graphic_locks.set("noChangeAspect", "1")
        c_nv_new.append(graphic_locks)
        anchor.append(c_nv_new)

    anchor.append(deepcopy(graphic))
    drawing.replace(inline, anchor)


def configure_section(section, options: BuildOptions) -> None:
    section.page_width = Mm(options.page_width_mm)
    section.page_height = Mm(options.page_height_mm)
    section.top_margin = Mm(options.margins.top)
    section.bottom_margin = Mm(options.margins.bottom)
    section.left_margin = Mm(options.margins.left)
    section.right_margin = Mm(options.margins.right)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)
    section.different_first_page_header_footer = False


def build_document(image_path: Path, options: BuildOptions) -> Document:
    doc = Document()
    section = doc.sections[0]
    configure_section(section, options)

    color = RGBColor(*options.font_color)
    style = doc.styles["Normal"]
    style.font.name = options.font_name
    style.font.size = Pt(options.font_size_pt)
    style.font.color.rgb = color
    style._element.rPr.rFonts.set(qn("w:eastAsia"), options.font_name)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    add_page_background_picture(hp, image_path, options.page_width_mm, options.page_height_mm)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = ""

    if doc.paragraphs:
        body = doc.paragraphs[0]
        body.text = ""
    else:
        body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = body.add_run("")
    set_run_font(run, options.font_name, options.font_size_pt, color)

    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)

    return doc


def docx_to_dotx(docx_path: Path, dotx_path: Path) -> None:
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
        dotx_path, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = text.replace(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                )
                data = text.encode("utf-8")
            zout.writestr(item, data)


def options_from_detection(detection: Detection, margins: Optional[MarginsMm] = None) -> BuildOptions:
    return BuildOptions(
        margins=margins or detection.margins,
        page_width_mm=detection.page_width_mm,
        page_height_mm=detection.page_height_mm,
    )


def write_packages(
    image_path: Path,
    options: BuildOptions,
    out_dir: Path,
    basename: str,
) -> Tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{basename}.docx"
    dotx_path = out_dir / f"{basename}.dotx"
    doc = build_document(image_path, options)
    doc.save(docx_path)
    docx_to_dotx(docx_path, dotx_path)
    return docx_path, dotx_path
